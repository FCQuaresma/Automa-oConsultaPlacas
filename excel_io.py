import openpyxl

import config


def abrir():
    wb = openpyxl.load_workbook(config.EXCEL_PATH)
    ws = wb.active
    return wb, ws


def listar_placas(ws):
    placas = []
    linha = config.DATA_START_ROW
    while True:
        valor = ws.cell(row=linha, column=config.PLACA_COL).value
        if valor is None or str(valor).strip() == "":
            break
        placas.append((linha, str(valor).strip().upper()))
        linha += 1
    return placas


def precisa_consultar(ws, linha, site):
    modelo = ws.cell(row=linha, column=site["modelo_col"]).value
    valor = ws.cell(row=linha, column=site["valor_col"]).value
    return not (modelo and str(modelo).strip()) or not (valor and str(valor).strip())


def gravar_resultado(ws, linha, site, modelo, valor):
    ws.cell(row=linha, column=site["modelo_col"], value=modelo)
    ws.cell(row=linha, column=site["valor_col"], value=valor)


def salvar(wb):
    wb.save(config.EXCEL_PATH)
